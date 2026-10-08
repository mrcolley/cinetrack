# tests/test_data_service.py
# Unit tests for parse_imdb_csv, filter_ratings, and export_to_excel_buffer.
# Also includes integration tests: parse → filter → export pipeline.

import io

import openpyxl
import pandas as pd
import pytest

from src.services.data_service import export_to_excel_buffer, filter_ratings, parse_imdb_csv

# ---------------------------------------------------------------------------
# parse_imdb_csv
# ---------------------------------------------------------------------------


def test_parse_imdb_csv_valid_string(sample_csv_string):
    df = parse_imdb_csv(sample_csv_string)
    assert isinstance(df, pd.DataFrame)
    assert len(df) == 10
    assert "Your_Rating" in df.columns
    assert "Title" in df.columns
    assert "Const" in df.columns


def test_parse_imdb_csv_valid_stringio(sample_csv_stringio):
    df = parse_imdb_csv(sample_csv_stringio)
    assert len(df) == 10


def test_parse_imdb_csv_valid_bytesio(sample_csv_bytesio):
    df = parse_imdb_csv(sample_csv_bytesio)
    assert len(df) == 10


def test_parse_imdb_csv_minimal_required_columns(minimal_csv_string):
    df = parse_imdb_csv(minimal_csv_string)
    assert "Const" in df.columns
    assert "Your_Rating" in df.columns
    assert "Title" in df.columns
    assert len(df) == 2


def test_parse_imdb_csv_missing_required_column(missing_col_csv_string):
    with pytest.raises(ValueError, match="Title"):
        parse_imdb_csv(missing_col_csv_string)


def test_parse_imdb_csv_mixed_rating_types():
    csv = "Const,Your Rating,Title\ntt1,7,Movie A\ntt2,9,Movie B\n"
    df = parse_imdb_csv(csv)
    # Ratings stored as string "7" must be cast to integer.
    assert pd.api.types.is_integer_dtype(df["Your_Rating"]) or str(df["Your_Rating"].dtype) == "Int64"
    assert df["Your_Rating"].iloc[0] == 7


def test_parse_imdb_csv_malformed_imdb_rating():
    csv = (
        "Const,Your Rating,Title,IMDb Rating\n"
        "tt1,8,Movie A,N/A\n"
        "tt2,6,Movie B,7.5\n"
    )
    df = parse_imdb_csv(csv)
    assert pd.isna(df["IMDb_Rating"].iloc[0])
    assert df["IMDb_Rating"].iloc[1] == 7.5


def test_parse_imdb_csv_malformed_date():
    csv = (
        "Const,Your Rating,Title,Date Rated\n"
        "tt1,8,Movie A,not-a-date\n"
        "tt2,6,Movie B,2024-01-15\n"
    )
    df = parse_imdb_csv(csv)
    assert pd.isna(df["Date_Rated"].iloc[0])
    assert not pd.isna(df["Date_Rated"].iloc[1])


def test_parse_imdb_csv_column_rename(sample_csv_string):
    df = parse_imdb_csv(sample_csv_string)
    # Raw IMDb headers should be normalised.
    assert "Your Rating" not in df.columns
    assert "Your_Rating" in df.columns
    assert "IMDb Rating" not in df.columns
    assert "IMDb_Rating" in df.columns


# ---------------------------------------------------------------------------
# filter_ratings
# ---------------------------------------------------------------------------


def test_filter_ratings_min_boundary(sample_df):
    result = filter_ratings(sample_df, min_rating=5)
    assert (result["Your_Rating"] >= 5).all()


def test_filter_ratings_max_boundary(sample_df):
    result = filter_ratings(sample_df, max_rating=7)
    assert (result["Your_Rating"] <= 7).all()


def test_filter_ratings_full_range(sample_df):
    result = filter_ratings(sample_df, min_rating=1, max_rating=10)
    assert len(result) == len(sample_df)


def test_filter_ratings_empty_result(sample_df):
    result = filter_ratings(sample_df, min_rating=11, max_rating=10)
    assert len(result) == 0
    assert isinstance(result, pd.DataFrame)


def test_filter_ratings_search_case_insensitive(sample_df):
    result = filter_ratings(sample_df, search_query="shawshank")
    assert len(result) == 1
    assert "Shawshank" in result["Title"].iloc[0]


def test_filter_ratings_search_no_match(sample_df):
    result = filter_ratings(sample_df, search_query="zzznomatch")
    assert len(result) == 0


def test_filter_ratings_search_partial_match(sample_df):
    # "Title T" matches "Title Two", "Title Three", etc. but not "Shawshank".
    result = filter_ratings(sample_df, search_query="title t")
    assert len(result) > 0
    assert all("title" in t.lower() for t in result["Title"])


def test_filter_ratings_column_selection(sample_df):
    cols = ["Title", "Your_Rating"]
    result = filter_ratings(sample_df, selected_columns=cols)
    assert list(result.columns) == cols


def test_filter_ratings_column_selection_nonexistent_ignored(sample_df):
    # A column name that does not exist should be silently ignored.
    cols = ["Title", "Your_Rating", "NonExistentColumn"]
    result = filter_ratings(sample_df, selected_columns=cols)
    assert "NonExistentColumn" not in result.columns
    assert "Title" in result.columns


def test_filter_ratings_does_not_mutate_input(sample_df):
    original_len = len(sample_df)
    original_cols = list(sample_df.columns)
    filter_ratings(sample_df, min_rating=5, search_query="x", selected_columns=["Title"])
    assert len(sample_df) == original_len
    assert list(sample_df.columns) == original_cols


def test_filter_ratings_empty_dataframe():
    empty = pd.DataFrame(columns=["Your_Rating", "Title"])
    result = filter_ratings(empty, min_rating=1, max_rating=10)
    assert isinstance(result, pd.DataFrame)
    assert len(result) == 0


def test_filter_ratings_single_row(sample_df):
    single = sample_df.iloc[:1].copy()
    result = filter_ratings(single, min_rating=1, max_rating=10)
    assert len(result) == 1


# ---------------------------------------------------------------------------
# export_to_excel_buffer
# ---------------------------------------------------------------------------


def test_export_to_excel_buffer_readable(sample_df):
    buf = export_to_excel_buffer(sample_df)
    wb = openpyxl.load_workbook(buf)
    assert wb is not None


def test_export_to_excel_buffer_sheet_name(sample_df):
    buf = export_to_excel_buffer(sample_df)
    wb = openpyxl.load_workbook(buf)
    assert "IMDb Ratings" in wb.sheetnames


def test_export_to_excel_buffer_custom_sheet_name(sample_df):
    buf = export_to_excel_buffer(sample_df, sheet_name="My Ratings")
    wb = openpyxl.load_workbook(buf)
    assert "My Ratings" in wb.sheetnames


def test_export_to_excel_buffer_stream_position_zero(sample_df):
    buf = export_to_excel_buffer(sample_df)
    assert buf.tell() == 0


def test_export_to_excel_buffer_row_count(sample_df):
    buf = export_to_excel_buffer(sample_df)
    wb = openpyxl.load_workbook(buf)
    ws = wb["IMDb Ratings"]
    # +1 for header row.
    assert ws.max_row == len(sample_df) + 1


# ---------------------------------------------------------------------------
# Integration: parse → filter → export pipeline
# ---------------------------------------------------------------------------


def test_integration_parse_then_filter(sample_csv_string):
    """Normalised column names from parse_imdb_csv are consumed by filter_ratings."""
    df = parse_imdb_csv(sample_csv_string)
    result = filter_ratings(df, min_rating=8, max_rating=10)
    assert len(result) > 0
    assert (result["Your_Rating"] >= 8).all()


def test_integration_filter_then_export(sample_df):
    """filter_ratings result feeds cleanly into export_to_excel_buffer."""
    filtered = filter_ratings(sample_df, min_rating=7, max_rating=10)
    buf = export_to_excel_buffer(filtered)
    wb = openpyxl.load_workbook(buf)
    ws = wb["IMDb Ratings"]
    assert ws.max_row == len(filtered) + 1  # data rows + header


def test_integration_parse_filter_export_full_pipeline(sample_csv_string):
    """Full parse → filter → export chain produces a valid, readable workbook."""
    df = parse_imdb_csv(sample_csv_string)
    filtered = filter_ratings(df, min_rating=5, search_query="title", selected_columns=["Title", "Your_Rating"])
    buf = export_to_excel_buffer(filtered)
    wb = openpyxl.load_workbook(buf)
    ws = wb["IMDb Ratings"]
    # Header row columns should match selected columns.
    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    assert "Title" in headers
    assert "Your_Rating" in headers
